import math
import os
import time


def confirm(question):
    return input(question + ' (y/n) ').strip().lower() == 'y'


def startfile(file):
    try:
        os.startfile(file)
    except AttributeError:
        os.system('open ' + file)


def demo_cocos2d():
    import cocos

    class Cocos2dDemo(cocos.layer.Layer):
        def __init__(self):
            super(Cocos2dDemo, self).__init__()
            label = cocos.text.Label('cocos2d OK',
                                     font_name='Arial', font_size=32,
                                     anchor_x='center', anchor_y='center')
            label.position = (320, 240)
            self.add(label)
            scale = cocos.actions.ScaleBy(1.2, duration=0.2)
            label.do(cocos.actions.Repeat(
                scale + cocos.actions.Reverse(scale)))

            def end_demo():
                cocos.director.director.pop()
                cocos.director.director.window.close()
            self.do(cocos.actions.Delay(2.0) +
                    cocos.actions.CallFunc(end_demo))
    cocos.director.director.init()
    cocos.director.director.run(cocos.scene.Scene(Cocos2dDemo()))
    return confirm('Did the message "cocos2d OK" appear?')


def demo_pygame():
    import pygame
    import pygame.locals
    pygame.init()
    screen = pygame.display.set_mode((640, 480))
    screen_rect = screen.get_rect()
    font = pygame.font.Font(pygame.font.match_font('Arial'), 32)
    text = font.render('pygame OK', True, (255, 255, 255))
    text_rect = text.get_rect()
    quit = False
    ticks_start = pygame.time.get_ticks()
    while True:
        for event in pygame.event.get():
            if event.type == pygame.locals.QUIT:
                quit = True
                break
        ticks = pygame.time.get_ticks()
        if quit or (ticks - ticks_start) >= 2000:
            break
        scale = 1.0 + 0.2 * (200 - abs(ticks % 400 - 200)) / 400
        screen.fill((0, 0, 0))
        scaled_text = pygame.transform.scale(text,
                                             (int(text_rect.width * scale),
                                              int(text_rect.height * scale)))
        scaled_text_rect = scaled_text.get_rect()
        scaled_text_rect.centerx = screen_rect.centerx
        scaled_text_rect.centery = screen_rect.centery
        screen.blit(scaled_text, scaled_text_rect)
        pygame.display.flip()
    pygame.quit()
    # Workaround for pygame window not closing
    pygame.init()
    pygame.quit()
    return confirm('Did the message "pygame OK" appear?')


def demo_numpy_matplotlib():
    import matplotlib.pyplot
    import numpy
    x = numpy.arange(0.0, 2.0 * math.pi, 0.1)
    y = numpy.sin(x)
    matplotlib.pyplot.plot(x, y)
    matplotlib.pyplot.show(block=False)
    time.sleep(2.0)
    matplotlib.pyplot.close('all')
    return confirm('Did a plot of sin(x) appear?')


def demo_openpyxl():
    import openpyxl
    workbook = openpyxl.Workbook()
    cell = workbook.active['A1']
    cell.value = 'openpyxl OK'
    cell.font = openpyxl.styles.Font(name='Arial', size=32)
    workbook.save('demo.xlsx')
    startfile('demo.xlsx')
    time.sleep(2.0)
    return confirm('Did a spreadsheet with the message "openpyxl OK" appear?')


def demo_python_docx():
    import docx
    document = docx.Document()
    paragraph = document.add_paragraph('python-docx OK')
    font = paragraph.style.font
    font.name = 'Arial'
    font.size = docx.shared.Pt(32)
    document.save('demo.docx')
    startfile('demo.docx')
    time.sleep(2.0)
    return confirm('Did a document with the message "python-docx OK" appear?')


def demo_colorama():
    import colorama
    colorama.init()
    output = 'colorama OK'
    colored_output = ''
    colors = [
        colorama.Fore.RED, colorama.Fore.GREEN, colorama.Fore.YELLOW,
        colorama.Fore.BLUE, colorama.Fore.MAGENTA, colorama.Fore.CYAN
    ]
    for index in range(len(output)):
        i = index % len(colors)
        colored_output += colors[i] + output[index]
    colored_output += colorama.Style.RESET_ALL
    print(colored_output)
    return confirm('Is the message "colorama OK" above multi-coloured?')


def demo_others():
    try:
        import pgzero
        import jupyter
        import pandas
        import PyPDF2
        import idlexlib
        import pyfirmata
        import PIL
    except:
        return False
    else:
        return True


def main():
    demos = [
        demo_cocos2d, demo_pygame, demo_numpy_matplotlib,
        demo_openpyxl, demo_python_docx, demo_colorama,
        demo_others
    ]
    for d in demos:
        if not d():
            print('DEMO FAILED!')
            input('(Hit enter to quit)\n')
            return
    print('ALL DEMOS SUCCEEDED!')
    input('(Hit enter to quit)\n')


if __name__ == "__main__":
    main()
